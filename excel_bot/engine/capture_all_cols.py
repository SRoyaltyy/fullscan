"""Lean A–JO dump from Yahoo/rows (never Excel --from-cache).

Each ticker: chosen tile(s), rows 2–145, all 275 columns (value + fill).
~1 s/ticker/tile. This is the full workbook surface, not A–O.

  python engine/capture_all_cols.py --tickers AAPL
  python engine/capture_all_cols.py --all-rows --workers 4
  python engine/capture_all_cols.py --all-rows --tile prev --workers 4
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import date
from functools import partial

from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from backtest import pick_anchors, seed_anchor  # noqa: E402
from colors import ColorEngine  # noqa: E402
from evaluator import Evaluator  # noqa: E402
from stockhistory import serial  # noqa: E402
from xlrt import is_arr  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ROWS_DIR = os.environ.get("ROWS_DIR", os.path.join(ROOT, "data", "rows"))
MODEL = os.path.join(HERE, "model.json")
SPLIT_PATH = os.path.join(HERE, "holdout_split.json")
ALL_COLS = 275
ROW_START, ROW_END = 2, 145
LAST_LETTER = get_column_letter(ALL_COLS)  # JO; formulas run through JL


def deser(rows):
    return [{**r, "date": date.fromisoformat(r["date"][:10])} for r in rows]


def stride_pick(xs, n):
    if n <= 0:
        return []
    if len(xs) <= n:
        return list(xs)
    step = len(xs) / n
    return [xs[int(i * step)] for i in range(n)]


def tickers_from_split(n_disc, n_hold):
    split = json.load(open(SPLIT_PATH))
    have = {fn[:-5] for fn in os.listdir(ROWS_DIR) if fn.endswith(".json")}
    disc = sorted(t for t in split["discovery"] if t in have)
    hold = sorted(t for t in split["holdout"] if t in have)
    return stride_pick(disc, n_disc), stride_pick(hold, n_hold)


def tickers_all_rows():
    split = json.load(open(SPLIT_PATH))
    have = sorted(fn[:-5] for fn in os.listdir(ROWS_DIR) if fn.endswith(".json"))
    dset, hset = set(split["discovery"]), set(split["holdout"])
    disc = [t for t in have if t in dset]
    hold = [t for t in have if t in hset]
    extra = [t for t in have if t not in dset and t not in hset]
    return disc, hold, extra


def dest_name(ticker, tile):
    if tile == "prev":
        return f"{ticker}__tprev.json"
    return f"{ticker}.json"


def dump_anchor(ticker, rows, anchor):
    ev = Evaluator(MODEL, today=serial(anchor))
    ev.seed(seed_anchor(ev, ticker, rows, anchor))
    ce = ColorEngine(ev, MODEL)
    days = []
    for r in range(ROW_START, ROW_END + 1):
        a = ev.get_cell(f"A{r}")
        if not isinstance(a, (int, float)) or a < 30000:
            continue
        cells = {}
        for c in range(1, ALL_COLS + 1):
            letter = get_column_letter(c)
            coord = f"{letter}{r}"
            v = ev.get_cell(coord)
            if is_arr(v):
                v = v[0][0]
            if hasattr(v, "code"):
                val, err = None, v.code
            else:
                val, err = v, None
            fill = ce.fill_for(coord)
            if val is None and not fill and not err:
                continue
            rec = {}
            if val is not None:
                rec["v"] = val
            if fill:
                rec["f"] = fill
            if err:
                rec["e"] = err
            cells[letter] = rec
        ohlc = {}
        for name, col in (("close", "B"), ("open", "C"), ("high", "D"),
                          ("low", "E"), ("volume", "F")):
            rec = cells.get(col) or {}
            ohlc[name] = rec.get("v") if isinstance(rec.get("v"), (int, float)) else None
        days.append({"date": int(a), **ohlc, "cells": cells})
    return days


def capture_ticker(ticker, outdir, tile="latest"):
    src = os.path.join(ROWS_DIR, f"{ticker}.json")
    if not os.path.exists(src):
        return ticker, 0, 0.0, "no_rows"
    t0 = time.time()
    rows = deser(json.load(open(src)))
    if len(rows) < 40:
        return ticker, 0, time.time() - t0, "short_history"
    anchors = pick_anchors(rows, rows[0]["date"], rows[-1]["date"])
    if not anchors:
        return ticker, 0, time.time() - t0, "no_anchors"
    if tile == "latest":
        chosen = [anchors[-1]]
    elif tile == "prev":
        if len(anchors) < 2:
            return ticker, 0, time.time() - t0, "no_prev_tile"
        chosen = [anchors[-2]]
    elif tile == "all":
        chosen = list(anchors)
    else:
        return ticker, 0, time.time() - t0, f"bad_tile:{tile}"
    by_date = {}
    for anchor in chosen:
        for d in dump_anchor(ticker, rows, anchor):
            by_date[d["date"]] = d
    days = [by_date[k] for k in sorted(by_date)]
    os.makedirs(outdir, exist_ok=True)
    json.dump({"ticker": ticker, "anchor": str(chosen[-1]),
               "anchors": [str(a) for a in chosen],
               "tile": tile, "all_cols": True,
               "n_cols": ALL_COLS, "last_letter": LAST_LETTER,
               "source": "rows_cache",
               "seed": "yahoo_rows_cache", "days": days},
              open(os.path.join(outdir, dest_name(ticker, tile)), "w"),
              default=str)
    return ticker, len(days), time.time() - t0, None


def main():
    os.chdir(ROOT)
    ap = argparse.ArgumentParser()
    ap.add_argument("--tickers", nargs="*", default=[])
    ap.add_argument("--from-split", action="store_true")
    ap.add_argument("--all-rows", action="store_true")
    ap.add_argument("--n-disc", type=int, default=200)
    ap.add_argument("--n-hold", type=int, default=120)
    ap.add_argument("--out", default="research/all_cols_grids")
    ap.add_argument("--skip-existing", action="store_true", default=True)
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--tile", choices=("latest", "prev", "all"), default="latest",
                    help="latest = newest 140-ish days; prev = prior STEP tile "
                         "(writes TICKER__tprev.json); all = stitch every anchor")
    args = ap.parse_args()
    if args.force:
        args.skip_existing = False
    extra = []
    disc = hold = []
    if args.all_rows:
        disc, hold, extra = tickers_all_rows()
        tickers = disc + hold + extra
        print(f"all-rows disc={len(disc)} hold={len(hold)} extra={len(extra)} "
              f"n={len(tickers)}", flush=True)
    elif args.from_split:
        disc, hold = tickers_from_split(args.n_disc, args.n_hold)
        tickers = disc + hold
        print(f"from-split disc={len(disc)} hold={len(hold)}", flush=True)
    else:
        tickers = args.tickers
    if not tickers:
        raise SystemExit("no tickers")
    outdir = args.out if os.path.isabs(args.out) else os.path.join(ROOT, args.out)
    os.makedirs(outdir, exist_ok=True)
    print(f"capture_all_cols n={len(tickers)} workers={args.workers} "
          f"tile={args.tile} -> {outdir}", flush=True)
    times, ok, err, skipped, todo = [], 0, 0, 0, []
    for t in tickers:
        dest = os.path.join(outdir, dest_name(t, args.tile))
        if args.skip_existing and os.path.exists(dest):
            skipped += 1
        else:
            todo.append(t)
    if skipped:
        print(f"  skip-existing {skipped}", flush=True)

    def _record(name, n, sec, e):
        nonlocal ok, err
        times.append(sec)
        if e:
            err += 1
            print(f"  FAIL {name} {sec:.1f}s {e}", flush=True)
        else:
            ok += 1
            if ok % 25 == 0 or ok + err == len(todo):
                print(f"  ok {name} days={n} {sec:.1f}s "
                      f"({ok+err}/{len(todo)} new, skip={skipped})", flush=True)

    worker = partial(capture_ticker, outdir=outdir, tile=args.tile)
    if args.workers <= 1 or len(todo) <= 1:
        for t in todo:
            _record(*worker(t))
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            futs = [pool.submit(worker, t) for t in todo]
            for fut in as_completed(futs):
                _record(*fut.result())
    measured = [s for s in times if s > 0]
    meta = {
        "n_ok": ok, "n_err": err, "n_skipped": skipped,
        "n_target": len(tickers),
        "n_on_disk": sum(1 for fn in os.listdir(outdir)
                         if fn.endswith(".json") and not fn.startswith("_")),
        "all_rows": bool(args.all_rows),
        "seconds_mean": (sum(measured) / len(measured)) if measured else None,
        "seed": "yahoo_rows_cache",
        "from_cache_flag": False,
        "cols": f"A..{LAST_LETTER}",
        "n_cols": ALL_COLS,
        "tile": args.tile,
    }
    json.dump(meta, open(os.path.join(outdir, "_meta.json"), "w"), indent=1)
    print(f"DONE ok={ok} skip={skipped} err={err} s/ticker={meta['seconds_mean']}")


if __name__ == "__main__":
    main()
